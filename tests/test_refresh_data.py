import copy,datetime as dt,importlib.util,json,pathlib,sys,tempfile,unittest
from unittest.mock import patch
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
import refresh_data as R

class Refresh(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.p=pathlib.Path(self.tmp.name)
        self.master=self.p/'master.xlsx';w=openpyxl.Workbook();s=w.active;s.title='Interest Rate Decisions'
        self.columns=['Effective Date']+list(dict.fromkeys(v['col'] for k,v in R.INDICATORS.items() if k!='crb_index'))
        s.append(self.columns);s.append([dt.datetime(2026,9,1)]+[None]*(len(self.columns)-1));w.save(self.master);w.close()
        core,_,_=R.read_core(self.master.read_bytes(),s.title,'Effective Date')
        self.base={}
        for slug,v in core.items():
            self.base[slug]=copy.deepcopy(v);self.base[slug]['latestObservationMonth']=None
        self.baseline=self.p/'approved.json';self.baseline.write_bytes(R.encode(self.base))
        self.extra=self.p/'extra.csv';self.extra.write_text('unused')
        self.meta=self.p/'meta.json';self.meta.write_text('{}')
        self.c={'workbook':{'path':str(self.master),'sha256':R.digest(self.master.read_bytes()),'logical_path':'Dropbox test master','revision':'test-revision'},'baseline':str(self.baseline),'baseline_sha256':R.digest(self.baseline.read_bytes()),'additional':str(self.extra),'additional_sha256':R.digest(self.extra.read_bytes()),'additional_metadata':str(self.meta),'additional_metadata_sha256':R.digest(self.meta.read_bytes()),'output_root':str(self.p/'staging'),'approvals':[]}
    def change(self,slug,value,date=None):
        w=openpyxl.load_workbook(self.master);s=w.active;s.cell(2,self.columns.index(R.INDICATORS[slug]['col'])+1).value=value
        if date is not None:s.cell(2,1).value=date
        w.save(self.master);w.close();self.c['workbook']['sha256']=R.digest(self.master.read_bytes())
    def run_refresh(self,mode='on-demand'):
        # Synthetic extra merge retains schema; real 59-series merge tested independently.
        with patch.object(R,'merge',side_effect=lambda core,*a:(dict((k,{**v,'latestObservationMonth':None}) for k,v in core.items()),{})):
            return R.run(self.c,mode)
    def test_explicit_workbook_acquisition_and_month_mapping(self):
        self.change('cpi_yoy',4.5);self.c['approvals']=[{'indicator':'cpi_yoy','month':'2026-09','old':None,'new':4.5}]
        result=self.run_refresh();d=json.loads((pathlib.Path(result['run'])/'data/indicators/cpi_yoy.json').read_text())
        self.assertEqual(d['series'],[{'date':'2026-09-01','value':4.5}]);self.assertEqual(result['staged_changes'],1)
    def test_unapproved_values_held_with_readable_audit(self):
        self.change('vacancy_rate',1);r=self.run_refresh();self.assertEqual(r['held_changes'],1)
        p=pathlib.Path(r['run']);self.assertIn('[held;', (p/'REPORT.md').read_text())
        self.assertIsNone(json.loads((p/'data/indicators/vacancy_rate.json').read_text())['series'][0]['value'])
    def test_historical_numeric_revisions_and_missing_are_held(self):
        self.base['cpi_yoy']['series'][0]['value']=4;self.base['cpi_yoy']['latestObservationMonth']='2026-09'
        self.baseline.write_bytes(R.encode(self.base));self.c['baseline_sha256']=R.digest(self.baseline.read_bytes())
        for value,kind in [(4.5,'historical_revision'),(None,'observed_to_missing')]:
            self.change('cpi_yoy',value);r=self.run_refresh();p=pathlib.Path(r['run'])
            self.assertEqual(json.loads((p/'changes.json').read_text())[0]['classification'],kind)
            self.assertEqual(json.loads((p/'data/indicators/cpi_yoy.json').read_text())['series'][0]['value'],4)
    def test_repeat_run_is_immutable_and_idempotent(self):
        a=self.run_refresh();b=self.run_refresh();self.assertEqual(a,b)
        self.assertEqual(len(list((self.p/'staging/runs').iterdir())),1)
    def test_dropbox_unavailable_preserves_last_candidate(self):
        a=self.run_refresh();p=pathlib.Path(a['run']);before={str(x):x.read_bytes() for x in p.rglob('*') if x.is_file()}
        self.c['workbook']['path']=str(self.p/'unavailable.xlsx')
        with self.assertRaisesRegex(ValueError,'unavailable'):self.run_refresh()
        self.assertEqual(before,{str(x):x.read_bytes() for x in p.rglob('*') if x.is_file()})
    def test_cloud_download_failure_has_no_silent_local_fallback(self):
        with patch.object(R.urllib.request,'urlopen',side_effect=OSError('private URL')):
            with self.assertRaisesRegex(ValueError,'no fallback'):R.acquire(self.c['workbook'],'https://example.test/private')
    def test_checksum_and_baseline_drift_fail(self):
        self.c['workbook']['sha256']='bad'
        with self.assertRaisesRegex(ValueError,'checksum'):self.run_refresh()
        self.c['workbook']['sha256']=R.digest(self.master.read_bytes());self.c['baseline_sha256']='bad'
        with self.assertRaisesRegex(ValueError,'baseline'):self.run_refresh()
    def test_duplicate_month_and_invalid_dates_fail(self):
        with self.assertRaises(ValueError):R.index([{'date':'2026-09-01','value':1},{'date':'2026-09-02','value':2}])
        for date in ['09/01/26','2026-02-30']:
            with self.assertRaises(ValueError):R.month(date)
    def test_historical_dates_metadata_and_ids_cannot_drift(self):
        for change in ['date','unit','id']:
            new=copy.deepcopy(self.base)
            if change=='date':new['cpi_yoy']['series'][0]['date']='2026-09-02'
            elif change=='unit':new['cpi_yoy']['unit']='different'
            else:del new['cpi_yoy']
            with self.assertRaises(ValueError):R.reconcile(self.base,new,[],lambda *a:{})
    def test_new_month_requires_approval_and_never_fills_other_months(self):
        new=copy.deepcopy(self.base);new['cpi_yoy']['series'].append({'date':'2026-10-01','value':4.6})
        candidate,audit=R.reconcile(self.base,new,[],lambda *a:{})
        self.assertEqual(len(candidate['cpi_yoy']['series']),1);self.assertEqual(audit[0]['classification'],'new_month')
    def test_missing_required_column_and_uncached_formula_fail(self):
        w=openpyxl.load_workbook(self.master);w.active.cell(1,2).value='wrong';w.save(self.master);w.close()
        with self.assertRaisesRegex(ValueError,'columns'):R.read_core(self.master.read_bytes(),'Interest Rate Decisions','Effective Date')
        w=openpyxl.load_workbook(self.master);w.active.cell(1,2).value=self.columns[1];w.active.cell(2,2).value='=1+1';w.save(self.master);w.close()
        with self.assertRaisesRegex(ValueError,'formula cache'):R.read_core(self.master.read_bytes(),'Interest Rate Decisions','Effective Date')
    def test_staging_cannot_target_checkout_or_input_directory(self):
        for out in [str(ROOT/'data'),str(self.p)]:
            self.c['output_root']=out
            with self.assertRaisesRegex(ValueError,'Unsafe'):self.run_refresh()
    def test_monthly_requires_verified_extraction(self):
        with self.assertRaisesRegex(ValueError,'extraction artifacts'):self.run_refresh('monthly')
    def automatic(self,value=1):
        staged=self.p/'automatic.csv';staged.write_text('month,excel_column,value,provider\n2026-09,AA,'+str(value)+',rba_xlsx\n')
        report=self.p/'report.json';report.write_bytes(R.encode({'latest_complete_month':'2026-09','series':{'AA':{'status':'extracted_pending_validation','count':40,'last':'2026-09'}}}))
        record=self.p/'run.json';record.write_bytes(R.encode({'status':'published','staged_sha256':R.digest(staged.read_bytes()),'validation_sha256':R.digest(report.read_bytes())}))
        self.c['automatic']={'staged':str(staged),'staged_sha256':R.digest(staged.read_bytes()),'report':str(report),'report_sha256':R.digest(report.read_bytes()),'run_record':str(record),'run_record_sha256':R.digest(record.read_bytes()),'mapping':{'AA':'aud_usd'},'max_lag_months':{'AA':0},'approved_providers':{'AA':'rba_xlsx'}}
    def test_monthly_shares_candidate_logic_and_holds_downloaded_values(self):
        self.automatic();a=self.run_refresh();b=self.run_refresh('monthly')
        self.assertEqual((pathlib.Path(a['run'])/'data/data.js').read_bytes(),(pathlib.Path(b['run'])/'data/data.js').read_bytes())
        self.assertEqual(b['held_changes'],1)
    def test_explicit_auto_fill_does_not_overwrite_manual_numeric_value(self):
        self.automatic();self.c['automatic']['approvals']=[{'indicator':'aud_usd','month':'2026-09','old':None,'new':1}]
        result=self.run_refresh('monthly');self.assertEqual(result['staged_changes'],1)
        self.base['aud_usd']['series'][0]['value']=0.7;self.base['aud_usd']['latestObservationMonth']='2026-09'
        self.baseline.write_bytes(R.encode(self.base));self.c['baseline_sha256']=R.digest(self.baseline.read_bytes());self.change('aud_usd',0.7)
        self.c['automatic']['approvals']=[{'indicator':'aud_usd','month':'2026-09','old':0.7,'new':1}]
        result=self.run_refresh('monthly');self.assertEqual(result['staged_changes'],0);self.assertEqual(result['held_changes'],1)
    def test_monthly_artifact_failure_and_excess_lag_block(self):
        self.automatic();self.c['automatic']['staged_sha256']='bad'
        with self.assertRaisesRegex(ValueError,'artifact changed'):self.run_refresh('monthly')
        self.automatic();report=pathlib.Path(self.c['automatic']['report']);v=json.loads(report.read_text());v['series']['AA']['last']='2026-07';report.write_bytes(R.encode(v));self.c['automatic']['report_sha256']=R.digest(report.read_bytes())
        record=pathlib.Path(self.c['automatic']['run_record']);v=json.loads(record.read_text());v['validation_sha256']=R.digest(report.read_bytes());record.write_bytes(R.encode(v));self.c['automatic']['run_record_sha256']=R.digest(record.read_bytes())
        with self.assertRaisesRegex(ValueError,'lag exceeded'):self.run_refresh('monthly')
    def test_lock_prevents_overlapping_runs(self):
        import fcntl
        p=pathlib.Path(self.c['output_root']);p.mkdir()
        with (p/'refresh.lock').open('a') as lock:
            fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
            with self.assertRaises(BlockingIOError):self.run_refresh()
    def test_snapshot_tampering_is_detected(self):
        self.run_refresh();p=next((self.p/'staging/snapshots').glob('*.xlsx'));p.chmod(0o600);p.write_bytes(b'bad')
        with self.assertRaisesRegex(ValueError,'Snapshot integrity'):self.run_refresh()

if __name__=='__main__':unittest.main()
