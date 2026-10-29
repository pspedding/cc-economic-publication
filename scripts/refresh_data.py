"""Staging-only unified refresh. No scheduler, source writes, Git writes or publication."""
import argparse, ast, copy, csv, datetime as dt, fcntl, hashlib, io, json, math
import operator, os, pathlib, re, shutil, tempfile, urllib.request
import openpyxl
from build_data import INDICATORS
from macro_data import merge, write_outputs

REPO=pathlib.Path(__file__).resolve().parents[1]
def load_config(path):
    path=pathlib.Path(path).resolve();c=json.loads(path.read_text())
    def resolve(v):return str((path.parent/pathlib.Path(v).expanduser()).resolve())
    for k in ['baseline','additional','additional_metadata','output_root']:c[k]=resolve(c[k])
    if c['workbook'].get('path'):c['workbook']['path']=resolve(c['workbook']['path'])
    c['protected_paths']=[resolve(p) for p in c.get('protected_paths',[])]
    if c.get('automatic'):
        for k in ['staged','report','run_record']:c['automatic'][k]=resolve(c['automatic'][k])
    return c
def encode(x):return (json.dumps(x,indent=2,allow_nan=False)+'\n').encode()
def digest(b):return hashlib.sha256(b).hexdigest()
def baseline_bytes(b):
    text=b.decode()
    if text.lstrip().startswith('{'):return json.loads(text)
    prefix='const CC_ECONOMIC_DATA = '
    if prefix not in text:raise ValueError('Unknown approved baseline format')
    return json.JSONDecoder().raw_decode(text.split(prefix,1)[1])[0]
def month(date):
    if isinstance(date,(dt.date,dt.datetime)):return date.strftime('%Y-%m-%d')
    if isinstance(date,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',date):
        dt.date.fromisoformat(date);return date
    raise ValueError('Unsupported observation date')
def index(series):
    result={}
    for p in series:
        date=month(p['date']);key=date[:7];v=p['value']
        if key in result:raise ValueError('Duplicate observation month')
        if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v)):raise ValueError('Invalid observation')
        result[key]=p
    return result
def formula_check(formulas,values):
    ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
    def evaluate(sheet,coord,seen=frozenset()):
        key=(sheet,coord)
        if key in seen:raise ValueError('Circular formula')
        value=formulas[sheet][coord].value
        if not isinstance(value,str) or not value.startswith('='):
            if value is not None and (isinstance(value,bool) or not isinstance(value,(int,float))):raise ValueError('Nonnumeric formula dependency')
            return 0 if value is None else value
        def ev(n):
            if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)):return n.value
            if isinstance(n,ast.Name) and re.fullmatch('[A-Z]+[1-9][0-9]*',n.id):return evaluate(sheet,n.id,seen|{key})
            if isinstance(n,ast.BinOp) and type(n.op) in ops:return ops[type(n.op)](ev(n.left),ev(n.right))
            if isinstance(n,ast.UnaryOp) and isinstance(n.op,ast.USub):return -ev(n.operand)
            raise ValueError('Unsupported formula; Excel verification required')
        return ev(ast.parse(value[1:],mode='eval').body)
    count=0
    for sheet in formulas:
        for row in sheet:
            for c in row:
                if c.data_type=='f':
                    count+=1;cache=values[sheet.title][c.coordinate].value;computed=evaluate(sheet.title,c.coordinate)
                    if not isinstance(cache,(int,float)) or not math.isclose(cache,computed,rel_tol=1e-10,abs_tol=1e-10):raise ValueError('Missing/stale formula cache')
    return count
def read_core(raw,sheet,date_column,active_slugs=None):
    f=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    w=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    try:
        caches=formula_check(f,w);rows=list(w[sheet].iter_rows());headers=[c.value for c in rows[0]]
        names=[h for h in headers if h is not None]
        if len(names)!=len(set(names)):raise ValueError('Duplicate workbook column')
        cols={v:i for i,v in enumerate(headers) if v is not None}
        active_slugs=(set(INDICATORS)-{'crb_index'}) if active_slugs is None else set(active_slugs)
        required={date_column}|{cfg['col'] for slug,cfg in INDICATORS.items() if slug in active_slugs}
        if required-set(cols):raise ValueError('Missing required workbook columns')
        data={};cells={}
        for slug,cfg in INDICATORS.items():
            if slug not in active_slugs:continue
            series=[]
            for row_number,row in enumerate(rows[1:],2):
                date=row[cols[date_column]].value
                if date is None:
                    orphan=row[cols[cfg['col']]].value
                    if orphan is not None and not (isinstance(orphan,str) and not orphan.strip()):raise ValueError('Observation without date')
                    continue
                date=month(date);cell=row[cols[cfg['col']]];value=cell.value
                if cell.data_type=='e':raise ValueError('Excel error in observation')
                if value is not None and (isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value)):raise ValueError('Invalid workbook value')
                series.append({'date':date,'value':value});cells[(slug,date[:7])]=openpyxl.utils.get_column_letter(cols[cfg['col']]+1)+str(row_number)
            series.sort(key=lambda p:p['date']);index(series)
            fields=['name','unit','freq','section','group','source','format','description','contextSeries','contextNarrative']
            data[slug]={k:copy.deepcopy(cfg[k]) for k in fields}
            data[slug].update(slug=slug,series=series)
        return data,cells,caches
    finally:f.close();w.close()
def acquire(source,url=None):
    try:
        if url:
            if not url.startswith('https://'):raise ValueError('HTTPS required')
            with urllib.request.urlopen(url,timeout=30) as r:raw=r.read(20*1024*1024+1)
        else:raw=pathlib.Path(source['path']).expanduser().read_bytes()
    except Exception:raise ValueError('Workbook unavailable; no fallback or publication') from None
    if len(raw)>20*1024*1024 or digest(raw)!=source['sha256']:raise ValueError('Workbook checksum mismatch')
    return raw
def snapshot(root,raw,suffix):
    path=root/(digest(raw)+suffix)
    if path.exists():
        if path.read_bytes()!=raw:raise ValueError('Snapshot integrity failure')
    else:
        with path.open('xb') as f:f.write(raw)
        path.chmod(0o400)
    return path.name
def reconcile(old,proposed,approvals,provenance):
    if set(old)!=set(proposed):raise ValueError('Unexpected indicator ID change')
    result=copy.deepcopy(old);audit=[]
    for slug,new in proposed.items():
        prior=old[slug];a=index(prior['series']);b=index(new['series'])
        # Source authority/schema changes cannot be authorized as numerical fills.
        ignored={'series','latestObservationMonth'}
        if {k:v for k,v in prior.items() if k not in ignored}!={k:v for k,v in new.items() if k not in ignored}:raise ValueError('Unexpected metadata/authority change: '+slug)
        selected=copy.deepcopy(a)
        for key in sorted(set(a)|set(b)):
            op=a.get(key);np=b.get(key)
            if op==np:continue
            # Existing source dates cannot drift, even with a value approval.
            if op and np and op['date']!=np['date']:raise ValueError('Historical date change')
            ov=op['value'] if op else None;nv=np['value'] if np else None
            rule={'indicator':slug,'month':key,'old':ov,'new':nv}
            approved=rule in approvals and np is not None
            kind='new_month' if op is None else 'removed_month' if np is None else 'missing_to_observed' if ov is None else 'observed_to_missing' if nv is None else 'historical_revision'
            audit.append({**rule,'classification':kind,'status':'approved_for_staging' if approved else 'held',
                'old_date':op['date'] if op else None,'new_date':np['date'] if np else None,**provenance(slug,key)})
            if approved:selected[key]=copy.deepcopy(np)
        result[slug]['series']=[selected[k] for k in sorted(selected)]
        have=[p['date'][:7] for p in result[slug]['series'] if p['value'] is not None]
        result[slug]['latestObservationMonth']=max(have) if have else None
    return result,audit
def run(config,mode='on-demand',workbook_url=None):
    """Both modes stage only. Monthly consumes an existing validated extraction run."""
    if mode not in {'on-demand','monthly'}:raise ValueError('Unknown mode')
    out=pathlib.Path(config['output_root']).expanduser().resolve()
    # Never write into this checkout, a source location, or any configured protection.
    protected=[REPO]+[pathlib.Path(p).expanduser().resolve() for p in config.get('protected_paths',[])]
    inputs=[pathlib.Path(config[k]).expanduser().resolve() for k in ['baseline','additional','additional_metadata']]
    if config['workbook'].get('path'):inputs.append(pathlib.Path(config['workbook']['path']).expanduser().resolve())
    if (out/'.git').exists() or any(out==p or out.is_relative_to(p) for p in protected) or any(p==out or p.is_relative_to(out) for p in inputs):raise ValueError('Unsafe staging output location')
    out.mkdir(parents=True,exist_ok=True)
    with (out/'refresh.lock').open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        raw=acquire(config['workbook'],workbook_url)
        base_raw=pathlib.Path(config['baseline']).read_bytes()
        if digest(base_raw)!=config['baseline_sha256']:raise ValueError('Approved baseline changed')
        old=baseline_bytes(base_raw)
        core,cells,caches=read_core(raw,config.get('sheet','Interest Rate Decisions'),config.get('date_column','Effective Date'),set(old)&set(INDICATORS)-{'crb_index'})
        csvraw=pathlib.Path(config['additional']).read_bytes();metaraw=pathlib.Path(config['additional_metadata']).read_bytes()
        if digest(csvraw)!=config['additional_sha256'] or digest(metaraw)!=config['additional_metadata_sha256']:raise ValueError('Additional source changed; review required')
        # Merge must read exactly captured inputs, not files that can change mid-run.
        with tempfile.TemporaryDirectory() as t:
            csvpath=pathlib.Path(t)/'additional.csv';mpath=pathlib.Path(t)/'metadata.json'
            csvpath.write_bytes(csvraw);mpath.write_bytes(metaraw)
            proposed,excluded=merge(core,csvpath,mpath)
        approvals=config.get('approvals',[])
        prov=lambda slug,key:{'source':config['workbook']['logical_path'] if slug in core else 'validated additional indicators',
          'source_sha256':digest(raw) if slug in core else digest(csvraw),'source_revision':config['workbook']['revision'] if slug in core else digest(csvraw),
          'sheet':config.get('sheet','Interest Rate Decisions') if slug in core else '', 'cell':cells.get((slug,key),'')}
        candidate,audit=reconcile(old,proposed,approvals,prov)
        automatic=None
        if mode=='monthly':
            automatic=config.get('automatic')
            if not automatic:raise ValueError('Monthly mode requires validated extraction artifacts')
            staged=pathlib.Path(automatic['staged']).read_bytes();report=pathlib.Path(automatic['report']).read_bytes()
            if digest(staged)!=automatic['staged_sha256'] or digest(report)!=automatic['report_sha256']:raise ValueError('Extraction artifact changed')
            status=json.loads(report);record_raw=pathlib.Path(automatic['run_record']).read_bytes();record=json.loads(record_raw)
            if digest(record_raw)!=automatic['run_record_sha256'] or record.get('status') not in {'published','dry_run_passed','no_changes'}:raise ValueError('Extraction run not validated')
            if record.get('staged_sha256')!=digest(staged) or record.get('validation_sha256')!=digest(report):raise ValueError('Extraction run artifacts disagree')
            for col in automatic['mapping']:
                v=status['series'].get(col,{})
                if v.get('status')!='extracted_pending_validation' or v.get('count',0)<36:raise ValueError('Automatic source unavailable')
                last=v['last'];month(last+'-01');complete=status['latest_complete_month'];month(complete+'-01')
                serial=lambda x:int(x[:4])*12+int(x[5:])
                if not 0<=serial(complete)-serial(last)<=automatic['max_lag_months'][col]:raise ValueError('Automatic source lag exceeded')
            # Compare external observations only. Never let them overwrite manual cells.
            mapping=automatic['mapping'];seen=set()
            for p in csv.DictReader(io.StringIO(staged.decode())):
                slug=mapping.get(p['excel_column'])
                if not slug:continue
                key=p['month'];month(key+'-01');pair=(slug,key)
                if pair in seen:raise ValueError('Duplicate automatic month')
                seen.add(pair);value=float(p['value'])
                if not math.isfinite(value):raise ValueError('Nonfinite automatic observation')
                if slug not in old:raise ValueError('Unknown automatic indicator')
                prior=index(candidate[slug]['series']).get(key);ov=prior['value'] if prior else None
                if value!=ov:
                    rule={'indicator':slug,'month':key,'old':ov,'new':value}
                    approved=ov is None and rule in automatic.get('approvals',[]) and p['provider']==automatic.get('approved_providers',{}).get(p['excel_column'])
                    # 10Y must pass the RBA F2.1 overlap gate, never bypass via FRED.
                    if slug=='bond_yield_10yr':approved=False
                    audit.append({**rule,'classification':'automatic_proposal','status':'approved_for_staging' if approved else 'held',
                      'source':p['provider'],'source_sha256':digest(staged),'source_revision':digest(report),'sheet':'','cell':'','old_date':prior['date'] if prior else None,'new_date':key+'-01'})
                    if approved:
                        points=index(candidate[slug]['series']);points[key]={'date':prior['date'] if prior else key+'-01','value':value}
                        candidate[slug]['series']=[points[k] for k in sorted(points)]
                        candidate[slug]['latestObservationMonth']=max(p['date'][:7] for p in points.values() if p['value'] is not None)
            automatic={'staged':staged,'report':report,'run_record':record_raw}
        manifest={'workbook_sha256':digest(raw),'workbook_revision':config['workbook']['revision'],'logical_source':config['workbook']['logical_path'],
          'baseline_sha256':digest(base_raw),'additional_sha256':digest(csvraw),'additional_metadata_sha256':digest(metaraw),
          'mode':mode,'formula_caches_checked':caches,'approvals':approvals,'automatic':{k:digest(v) for k,v in (automatic or {}).items()}}
        manifest['date_policy']={'sheet':config.get('sheet','Interest Rate Decisions'),'column':config.get('date_column','Effective Date'),'semantics':'preserve_source_dates; calendar-month join'}
        manifest['automatic_policy']={k:config['automatic'].get(k,{}) for k in ['mapping','max_lag_months','approvals','approved_providers']} if mode=='monthly' else None
        # Bind run identity to processing code and policy, not wall time.
        manifest['code_sha256']={p.name:digest(p.read_bytes()) for p in [pathlib.Path(__file__),REPO/'scripts/build_data.py',REPO/'scripts/macro_data.py']}
        runid=digest(encode(manifest));dest=out/'runs'/runid
        snapshots=out/'snapshots';snapshots.mkdir(exist_ok=True)
        for b,s in [(raw,'.xlsx'),(base_raw,'.json'),(csvraw,'.csv'),(metaraw,'.json')]+[(b,'.'+k) for k,b in (automatic or {}).items()]:snapshot(snapshots,b,s)
        with tempfile.TemporaryDirectory(dir=out) as t:
            temp=pathlib.Path(t);write_outputs(candidate,excluded,temp/'data')
            (temp/'manifest.json').write_bytes(encode(manifest));(temp/'changes.json').write_bytes(encode(audit))
            lines=['# Refresh change report','',f'Mode: {mode}. Staging only; no publication.',f'Indicators: {len(candidate)}. Formula caches checked: {caches}.','']
            lines.extend(f"- {p['indicator']} {p['month']}: {p['old']} -> {p['new']} [{p['status']}; {p['classification']}] ({p['source']}, {p.get('sheet','')}!{p.get('cell','')})" for p in audit)
            if not audit:lines.append('No observation changes.')
            (temp/'REPORT.md').write_text('\n'.join(lines)+'\n')
            hashes={str(p.relative_to(temp)):digest(p.read_bytes()) for p in temp.rglob('*') if p.is_file()}
            (temp/'hashes.json').write_bytes(encode(dict(sorted(hashes.items()))))
            dest.parent.mkdir(exist_ok=True)
            if dest.exists():
                existing={str(p.relative_to(dest)):p.read_bytes() for p in dest.rglob('*') if p.is_file()}
                new={str(p.relative_to(temp)):p.read_bytes() for p in temp.rglob('*') if p.is_file()}
                if existing!=new:raise ValueError('Immutable run mismatch')
            else:os.rename(temp,dest)
        return {'run':str(dest),'run_id':runid,'indicators':len(candidate),'files':len(list((dest/'data').rglob('*.json')))+1,
          'staged_changes':sum(p['status']=='approved_for_staging' for p in audit),'held_changes':sum(p['status']=='held' for p in audit),'publication':'not_permitted'}
def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--config',type=pathlib.Path,required=True)
    p.add_argument('--mode',choices=['on-demand','monthly'],default='on-demand');a=p.parse_args()
    try:result=run(load_config(a.config),a.mode,os.environ.get('DASHBOARD_WORKBOOK_URL'))
    except Exception as e:p.exit(1,'Refresh failed: '+type(e).__name__+'; no publication or fallback\n')
    print(json.dumps(result,indent=2))
if __name__=='__main__':main()
